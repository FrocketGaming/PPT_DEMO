"""Reads a deck YAML and produces a .pptx. All slide content comes from the config.

Two modes:
- no `deck.template`: slides are drawn from scratch with the built-in theme
- `deck.template: <name>`: slides use the brand template's layouts, placeholders,
  colours and fonts; regions are laid out inside the template's content area
"""

from __future__ import annotations

import copy
import sys
from dataclasses import replace
from datetime import datetime, timezone
from pathlib import Path

import yaml
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

from . import components
from .components import Context, add_rich_text, render
from .data import load_sources
from .layout import Box, compute_regions
from .template import Template, default_date, find_template
from .theme import Theme

SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.5
CONTENT = Box(MARGIN, 1.55, SLIDE_W - 2 * MARGIN, SLIDE_H - 1.55 - 0.6)
REGION_KEYS = ("left", "center", "right", "top", "bottom", "full")
SLIDE_TYPES = ("title", "section", "content", "closing")
SLIDE_KEYS = {"type", "title", "subtitle", "source", "notes", "skip", "gap", "regions",
              "font_size", "layout", "placeholders", "tag", "date", "background", "eyebrow",
              "number", *REGION_KEYS}


class ConfigError(Exception):
    pass


def warn(msg: str) -> None:
    print(f"warning: {msg}", file=sys.stderr)


def _merge_defaults(spec: dict, defaults: dict) -> dict:
    merged = copy.deepcopy(defaults.get(spec.get("type"), {}))
    merged.update(spec)
    return merged


def _regions(slide_cfg: dict) -> dict:
    regions = {k: slide_cfg[k] for k in REGION_KEYS if k in slide_cfg}
    regions.update(slide_cfg.get("regions", {}))
    return regions


def _render_regions(slide, regions: dict, area: Box, ctx: Context, defaults: dict,
                    gap: float) -> None:
    boxes = compute_regions(regions, area, gap=gap)
    for name, region in regions.items():
        try:
            specs = region["stack"] if "stack" in region else [region]
            weights = [s.get("weight", 1) for s in specs]
            for spec, box in zip(specs, boxes[name].split_rows(weights, gap=0.2)):
                render(slide, box, _merge_defaults(spec, defaults), ctx)
        except Exception as exc:
            raise ValueError(f"region '{name}': {exc}") from exc


def _small_text(slide, box: Box, text: str, theme: Theme, size: float = 10,
                color: str = "muted", align: str = "left") -> None:
    tf = components._textbox(slide, box)
    tf.paragraphs[0].alignment = components.ALIGN[align]
    add_rich_text(tf.paragraphs[0], text, theme, size, color=color)


# --------------------------------------------------------------------------- built-in look

def _background(slide, slide_cfg: dict, theme: Theme, default: str | None) -> Theme:
    """Fills the slide background; returns the theme to draw on it with."""
    color = slide_cfg.get("background", default)
    if color is None:
        return theme
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = theme.rgb(color)
    return theme.on_background(color)


def _title_block(slide, box: Box, slide_cfg: dict, theme: Theme, size: float) -> None:
    """Title text, with the optional `eyebrow` (a small uppercase kicker) just above it."""
    tf = components._textbox(slide, box, anchor=MSO_ANCHOR.BOTTOM)
    p = tf.paragraphs[0]
    if eyebrow := slide_cfg.get("eyebrow"):
        add_rich_text(p, str(eyebrow).upper(), theme, max(12, size * 0.4), color="accent",
                      bold=True)
        p.space_after = Pt(size * 0.15)
        p = tf.add_paragraph()
    add_rich_text(p, slide_cfg.get("title", ""), theme, size, color="primary", bold=True)


def _plain_title_slide(prs, slide_cfg: dict, theme: Theme, kind: str, deck: dict):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    base = theme
    theme = _background(slide, slide_cfg, theme, "light" if kind == "section" else "primary")
    bg = slide_cfg.get("background", "light" if kind == "section" else "primary")

    if kind == "section" and (number := slide_cfg.get("number")) is not None:
        tf = components._textbox(slide, Box(SLIDE_W / 2, 0.2, SLIDE_W / 2 - 0.3, SLIDE_H - 0.4),
                                 anchor=MSO_ANCHOR.MIDDLE)
        tf.paragraphs[0].alignment = components.ALIGN["right"]
        faint = base.mix(bg, "FFFFFF" if base.is_dark(bg) else "primary", 0.07)
        add_rich_text(tf.paragraphs[0], str(number), theme, 300, color=faint, bold=True)
    elif kind != "section" and deck.get("decor", True):
        for x, y, d, t, color in ((8.9, 0.9, 7.2, 0.16, base.mix(bg, "FFFFFF", 0.07)),
                                  (10.6, -1.3, 4.0, 0.035, base.hex("accent"))):
            ring = components._shape(slide, MSO_SHAPE.DONUT, Box(x, y, d, d), base.rgb(color))
            ring.adjustments[0] = t

    components._shape(slide, MSO_SHAPE.RECTANGLE, Box(MARGIN, 3.55, 1.2, 0.08), theme.rgb("accent"))
    _title_block(slide, Box(MARGIN, 0.8, SLIDE_W * 0.62, 2.65), slide_cfg, theme,
                 slide_cfg.get("font_size", 40))
    if sub := slide_cfg.get("subtitle"):
        _small_text(slide, Box(MARGIN, 3.85, SLIDE_W * 0.62, 1.5), sub, theme, 20)
    if kind == "title" and (author := deck.get("author")):
        _small_text(slide, Box(MARGIN, SLIDE_H - 1.0, 6, 0.4),
                    f"**{author}**   ·   {slide_cfg.get('date', deck.get('date') or default_date())}",
                    theme, 14)
    return slide


def _plain_content_slide(prs, slide_cfg: dict, ctx: Context, defaults: dict, number: int,
                         total: int, section: str | None, footer: dict):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    ctx = replace(ctx, theme=_background(slide, slide_cfg, ctx.theme, None))
    theme = ctx.theme

    _title_block(slide, Box(MARGIN, 0.2, SLIDE_W - 2 * MARGIN, 0.9), slide_cfg, theme, 28)
    if sub := slide_cfg.get("subtitle"):
        _small_text(slide, Box(MARGIN, 1.08, SLIDE_W - 2 * MARGIN, 0.4), sub, theme, 15)
    if src := slide_cfg.get("source"):
        _small_text(slide, Box(MARGIN, SLIDE_H - 0.45, SLIDE_W - 2 * MARGIN - 5, 0.3),
                    f"Source: {src}", theme)
    _small_text(slide, Box(SLIDE_W - MARGIN - 0.6, SLIDE_H - 0.45, 0.6, 0.3), str(number),
                theme, align="right")
    if footer.get("chapter") and (chapter := slide_cfg.get("tag", section)):
        _small_text(slide, Box(SLIDE_W - MARGIN - 4.7, SLIDE_H - 0.45, 4, 0.3), f"**{chapter}**",
                    theme, align="right")
    if footer.get("progress") and total:
        track = Box(0, SLIDE_H - 0.06, SLIDE_W, 0.06)
        components._shape(slide, MSO_SHAPE.RECTANGLE, track, theme.rgb("light"))
        components._shape(slide, MSO_SHAPE.RECTANGLE,
                          Box(0, track.y, SLIDE_W * number / total, track.h), theme.rgb("accent"))

    regions = _regions(slide_cfg)
    if not regions:
        raise ValueError("content slide has no regions (use left/right/top/bottom/center/full)")
    _render_regions(slide, regions, CONTENT, ctx, defaults, slide_cfg.get("gap", 0.3))
    return slide


# --------------------------------------------------------------------------- template look

TEXT_STYLE_KEYS = {"text", "font_size", "bold", "italic", "color", "align"}


def _fill_text(placeholder, value, theme: Theme, style: dict | None = None) -> None:
    """Writes text into a placeholder, keeping the template's fonts, sizes and bullets.

    `style` overrides only what it sets: font_size, bold, italic, color, align.
    """
    style = style or {}
    lines = value if isinstance(value, list) else str(value).split("\n")
    tf = placeholder.text_frame
    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        add_rich_text(p, line, theme, size=None, color=None)
        if "align" in style:
            p.alignment = components.ALIGN[style["align"]]
        for run in p.runs:
            if "font_size" in style:
                run.font.size = Pt(style["font_size"])
            if "bold" in style:
                run.font.bold = style["bold"]
            if "italic" in style:
                run.font.italic = style["italic"]
            if "color" in style:
                run.font.color.rgb = theme.rgb(style["color"])


def _fill_placeholders(slide, fills: dict, ctx: Context, defaults: dict) -> None:
    """`placeholders: {idx: text | [lines] | {text, font_size, ...} | component}`."""
    by_idx = {ph.placeholder_format.idx: ph for ph in slide.placeholders}
    for idx, value in fills.items():
        if (ph := by_idx.get(int(idx))) is None:
            raise ValueError(f"placeholder {idx} not on layout '{slide.slide_layout.name}'; "
                             f"available: {sorted(by_idx)}")
        if not isinstance(value, dict):
            _fill_text(ph, value, ctx.theme)
            continue
        if "type" not in value:  # styled text: {text: ..., font_size: 14, ...}
            if unknown := sorted(set(value) - TEXT_STYLE_KEYS):
                raise ValueError(f"placeholder {idx}: unknown key(s) {unknown}; styled text takes "
                                 f"{sorted(TEXT_STYLE_KEYS)} (or give a component `type`)")
            if "text" not in value:
                raise ValueError(f"placeholder {idx}: styled text needs a `text` key")
            _fill_text(ph, value["text"], ctx.theme, value)
            continue
        spec = _merge_defaults(value, defaults)
        if spec.get("type") == "image" and hasattr(ph, "insert_picture"):
            ph.insert_picture(str((ctx.base_dir / spec["path"]).resolve()))
            continue
        box = Box(*(Emu(v).inches for v in (ph.left, ph.top, ph.width, ph.height)))
        ph._element.getparent().remove(ph._element)
        try:
            render(slide, box, spec, ctx)
        except Exception as exc:
            raise ValueError(f"placeholder {idx}: {exc}") from exc


def _template_slide(tpl: Template, kind: str, slide_cfg: dict, ctx: Context, defaults: dict,
                    section: str | None, deck: dict):
    slide, roles = tpl.add_slide(kind, slide_cfg)
    values = {
        "title": slide_cfg.get("title"),
        "subtitle": slide_cfg.get("subtitle"),
        "tag": slide_cfg.get("tag", section if kind == "content" else None),
        "date": slide_cfg.get("date", deck.get("date") or default_date())
        if kind == "title" else None,
    }
    for role, value in values.items():
        if value is None:
            continue
        if role in roles:
            _fill_text(roles[role], value, ctx.theme)
        elif role == "title" or (role == "subtitle" and kind != "content"):
            warn(f"layout '{slide.slide_layout.name}' has no {role} placeholder; "
                 f"'{value}' not shown (map it in {tpl.path.stem}.yaml)")

    _fill_placeholders(slide, slide_cfg.get("placeholders", {}), ctx, defaults)

    regions = _regions(slide_cfg)
    if kind == "content" or regions:
        area = tpl.content_area(slide, roles.get("title"))
        if (sub := values["subtitle"]) and "subtitle" not in roles and kind == "content":
            _small_text(slide, Box(area.x, area.y, area.w, 0.4), sub, ctx.theme, 15)
            area = Box(area.x, area.y + 0.5, area.w, area.h - 0.5)
        if src := slide_cfg.get("source"):
            _small_text(slide, Box(area.x, area.y + area.h + 0.08, area.w, 0.28),
                        f"Source: {src}", ctx.theme)
        if regions:
            _render_regions(slide, regions, area, ctx, defaults, slide_cfg.get("gap", 0.3))
        elif not slide_cfg.get("placeholders"):
            raise ValueError("content slide has no regions or placeholders")

    tpl.remove_empty_placeholders(slide)
    return slide


# --------------------------------------------------------------------------- entry point

def _template_dirs(deck: dict, base_dir: Path) -> list[Path]:
    dirs = [base_dir / deck["templates_dir"]] if deck.get("templates_dir") else []
    return dirs + [base_dir / "templates", base_dir.parent / "templates", Path.cwd() / "templates"]


def build(config_path: str | Path, output: str | Path | None = None) -> Path:
    config_path = Path(config_path).resolve()
    try:
        cfg = yaml.safe_load(config_path.read_text(encoding="utf-8"))
    except yaml.YAMLError as exc:
        raise ConfigError(f"{config_path.name} is not valid YAML "
                          f"(check indentation: slide keys must line up):\n{exc}") from exc
    base_dir = config_path.parent
    deck = cfg.get("deck", {})

    tpl = None
    if name := deck.get("template"):
        tpl = Template(find_template(name, _template_dirs(deck, base_dir)))
        prs = tpl.prs
        theme = Theme.from_config({**tpl.theme, **(deck.get("theme") or {})})
    else:
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
        theme = Theme.from_config(deck.get("theme"))

    ctx = Context(theme=theme, data=load_sources(cfg.get("data", {}), base_dir), base_dir=base_dir)
    defaults = cfg.get("defaults", {})
    # set every identity field so nothing leaks through from the base/template file
    props, now = prs.core_properties, datetime.now(timezone.utc).replace(tzinfo=None)
    props.title = deck.get("title", "")
    props.author = props.last_modified_by = deck.get("author", "")
    props.created = props.modified = now
    props.revision = 1
    props.subject = props.keywords = props.comments = props.category = ""

    section = None
    footer = deck.get("footer") or {}
    total = sum(1 for s in cfg.get("slides", []) if not s.get("skip"))
    for i, slide_cfg in enumerate(cfg.get("slides", []), start=1):
        if slide_cfg.get("skip"):
            continue
        kind = slide_cfg.get("type", "content")
        try:
            if unknown := sorted(set(slide_cfg) - SLIDE_KEYS):
                raise ValueError(f"unknown key(s) {unknown}; valid: {sorted(SLIDE_KEYS)}")
            if kind not in SLIDE_TYPES:
                raise ValueError(f"unknown slide type '{kind}'; valid: {list(SLIDE_TYPES)}")
            if kind == "section":
                section = slide_cfg.get("title")
            if tpl:
                slide = _template_slide(tpl, kind, slide_cfg, ctx, defaults, section, deck)
            elif slide_cfg.get("layout") or slide_cfg.get("placeholders"):
                raise ValueError("'layout'/'placeholders' need a template (set deck.template)")
            elif kind in ("title", "section", "closing"):
                slide = _plain_title_slide(prs, slide_cfg, theme, kind, deck)
            else:
                slide = _plain_content_slide(prs, slide_cfg, ctx, defaults, len(prs.slides) + 1,
                                             total, section, footer)
        except Exception as exc:
            raise ConfigError(f"slide {i} ('{slide_cfg.get('title', '')}'): {exc}") from exc
        if notes := slide_cfg.get("notes"):
            slide.notes_slide.notes_text_frame.text = notes

    output = Path(output) if output else Path.cwd() / "out" / f"{config_path.stem}.pptx"
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    return output
