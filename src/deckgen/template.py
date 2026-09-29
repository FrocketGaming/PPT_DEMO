"""Building slides on top of a brand template (.pptx or .potx).

A deck opts in with `deck.template: <name>`; the name is looked up in a `templates/`
folder. An optional profile next to the template (`<name>.yaml`) says which layout to
use for each slide type and which placeholder index plays which role:

    slides:
      title:   {layout: "Presentation Title Cool Gray", title: 10, subtitle: 11, date: 12}
      content: {layout: "Content Slide – Title Only", title: 22, tag: 21}
    content_area: [x, y, w, h]      # inches; where regions are laid out
    theme: {palette: [...]}         # overrides colours read from the template

Without a profile, layouts and title placeholders are guessed from names and types.
"""

from __future__ import annotations

import copy
import io
import re
import zipfile
from datetime import date
from pathlib import Path

import yaml
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Emu

from .layout import Box

TEMPLATE_EXTS = (".pptx", ".potx")
POTX_TYPE = b"application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"
PPTX_TYPE = b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"

# layout-name keywords tried, in order, when no profile names a layout
LAYOUT_GUESSES = {
    "title": ["title slide", "presentation title"],
    "section": ["section header", "section divider", "section"],
    "content": ["title only"],
    "closing": ["thank you", "closing", "end"],
}
FOOTER_TYPES = (PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE)
TITLE_TYPES = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)


def _norm(name: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[–—]", "-", name)).strip().casefold()


def _inches(v) -> float:
    return Emu(v).inches


def find_template(name: str, search_dirs: list[Path]) -> Path:
    candidates = [Path(name)] if Path(name).suffix in TEMPLATE_EXTS else [
        Path(name + ext) for ext in TEMPLATE_EXTS]
    for d in search_dirs:
        for c in candidates:
            if (d / c).is_file():
                return (d / c).resolve()
    available = sorted({p.stem for d in search_dirs if d.is_dir()
                        for p in d.iterdir() if p.suffix in TEMPLATE_EXTS})
    raise FileNotFoundError(
        f"template '{name}' not found in {[str(d) for d in search_dirs]}; available: {available}")


def open_presentation(path: Path) -> Presentation:
    """python-pptx refuses .potx files, so relabel the package as a presentation in memory."""
    if path.suffix.lower() != ".potx":
        return Presentation(str(path))
    src, buf = zipfile.ZipFile(path), io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(POTX_TYPE, PPTX_TYPE)
            out.writestr(item, data)
    buf.seek(0)
    return Presentation(buf)


def theme_from_template(prs) -> dict:
    """Reads brand colours and body font from the template's theme."""
    theme_part = prs.slide_masters[0].part.part_related_by(RT.THEME)
    root = etree.fromstring(theme_part.blob)
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    colors = {}
    for el in root.find(".//a:clrScheme", ns):
        tag = etree.QName(el).localname
        srgb = el.find("a:srgbClr", ns)
        sys_ = el.find("a:sysClr", ns)
        colors[tag] = srgb.get("val") if srgb is not None else sys_.get("lastClr", "000000")

    def is_light(hex_: str) -> bool:
        r, g, b = (int(hex_[i:i + 2], 16) for i in (0, 2, 4))
        return 0.299 * r + 0.587 * g + 0.114 * b > 190

    accents = [colors[f"accent{i}"] for i in range(1, 7) if f"accent{i}" in colors]
    theme = {
        "primary": colors.get("dk2", colors.get("dk1")),
        "accent": colors.get("accent1"),
        "text": colors.get("dk1"),
        "light": colors.get("lt2", "F1F4F8"),
        "palette": [c for c in [colors.get("dk2"), *accents] if c and not is_light(c)],
    }
    font = root.find(".//a:minorFont/a:latin", ns)
    if font is not None and not font.get("typeface", "").startswith("+"):
        theme["font"] = font.get("typeface")
    return {k: v for k, v in theme.items() if v}


class Template:
    def __init__(self, path: Path):
        self.path = path
        self.prs = open_presentation(path)
        profile_path = path.with_suffix(".yaml")
        self.profile = (yaml.safe_load(profile_path.read_text(encoding="utf-8")) or {}
                        if profile_path.exists() else {})
        self.layouts = {_norm(l.name): l for l in self.prs.slide_layouts}
        self._remove_sample_slides()

    @property
    def theme(self) -> dict:
        return {**theme_from_template(self.prs), **self.profile.get("theme", {})}

    @property
    def size(self) -> tuple[float, float]:
        return _inches(self.prs.slide_width), _inches(self.prs.slide_height)

    def _remove_sample_slides(self) -> None:
        id_list = self.prs.slides._sldIdLst
        for sld_id in list(id_list):
            self.prs.part.drop_rel(sld_id.rId)
            id_list.remove(sld_id)

    # ------------------------------------------------------------------ layouts

    def layout(self, name: str):
        if (found := self.layouts.get(_norm(name))) is None:
            names = "\n  ".join(l.name for l in self.prs.slide_layouts)
            raise ValueError(f"layout '{name}' not in template {self.path.name}; layouts:\n  {names}")
        return found

    def role_spec(self, kind: str, override_layout: str | None) -> tuple[object, dict]:
        """Returns (layout, {role: placeholder idx}) for a slide type."""
        spec = dict(self.profile.get("slides", {}).get(kind, {}))
        if override_layout:
            spec["layout"] = override_layout
        if "layout" in spec:
            return self.layout(spec.pop("layout")), spec
        for keyword in LAYOUT_GUESSES.get(kind, []):
            for name, layout in self.layouts.items():
                if keyword in name:
                    return layout, spec
        raise ValueError(f"template {self.path.name} has no layout for '{kind}' slides; "
                         f"add one under `slides.{kind}.layout` in {self.path.stem}.yaml")

    # ------------------------------------------------------------------ slides

    def add_slide(self, kind: str, cfg: dict) -> tuple[object, dict]:
        """Adds a slide and returns (slide, {role: placeholder}) for the builder to fill."""
        layout, roles = self.role_spec(kind, cfg.get("layout"))
        slide = self.prs.slides.add_slide(layout)
        by_idx = {ph.placeholder_format.idx: ph for ph in slide.placeholders}
        found = {role: by_idx[idx] for role, idx in roles.items() if idx in by_idx}
        if "title" not in found:
            found.update(self._guess_title(slide))
        self._clone_footer_placeholders(slide, layout)
        return slide, found

    def _guess_title(self, slide) -> dict:
        phs = list(slide.placeholders)
        titles = [p for p in phs if p.placeholder_format.type in TITLE_TYPES]
        subs = [p for p in phs if p.placeholder_format.type == PP_PLACEHOLDER.SUBTITLE]
        if not titles:  # brand templates often use a wide body placeholder at the top
            width = self.prs.slide_width
            bodies = sorted((p for p in phs if p.placeholder_format.type == PP_PLACEHOLDER.BODY
                             and p.width > width * 0.4), key=lambda p: p.top)
            titles = bodies[:1]
        out = {"title": titles[0]} if titles else {}
        if subs:
            out["subtitle"] = subs[0]
        return out

    @staticmethod
    def _clone_footer_placeholders(slide, layout) -> None:
        """New slides don't inherit slide-number/footer/date placeholders; copy them over."""
        present = {ph.placeholder_format.type for ph in slide.placeholders}
        for ph in layout.placeholders:
            if ph.placeholder_format.type in FOOTER_TYPES and ph.placeholder_format.type not in present:
                slide.shapes._spTree.append(copy.deepcopy(ph._element))

    def content_area(self, slide, title_ph) -> Box:
        if area := self.profile.get("content_area"):
            return Box(*area)
        w, h = self.size
        left = _inches(title_ph.left) if title_ph is not None else 0.5
        top = _inches(title_ph.top + title_ph.height) + 0.15 if title_ph is not None else 1.5
        # stop above any master/layout decoration in the bottom quarter (footer bars, logos)
        decor_tops = [_inches(s.top) for s in [*slide.slide_layout.shapes,
                                               *slide.slide_layout.slide_master.shapes]
                      if not s.is_placeholder and s.top is not None
                      and _inches(s.top) > h * 0.75]
        bottom = min(decor_tops, default=h - 0.5) - 0.45
        return Box(left, top, w - 2 * left, bottom - top)

    @staticmethod
    def remove_empty_placeholders(slide) -> None:
        for ph in list(slide.placeholders):
            if ph.placeholder_format.type in FOOTER_TYPES:
                continue
            unfilled = type(ph).__name__ in ("PicturePlaceholder", "ChartPlaceholder",
                                             "TablePlaceholder")
            if unfilled or (ph.has_text_frame and not ph.text_frame.text.strip()):
                ph._element.getparent().remove(ph._element)


def default_date() -> str:
    return date.today().strftime("%B %Y").upper()


def describe(path: Path) -> str:
    """Human-readable list of layouts and placeholders, for writing profiles."""
    prs = open_presentation(path)
    lines = [f"{path.name}: {len(prs.slide_layouts)} layouts, "
             f"{_inches(prs.slide_width):.2f} x {_inches(prs.slide_height):.2f} in"]
    for layout in prs.slide_layouts:
        lines.append(f"\n{layout.name}")
        for ph in sorted(layout.placeholders, key=lambda p: (p.top or 0, p.left or 0)):
            pf = ph.placeholder_format
            if pf.type in FOOTER_TYPES:
                continue
            lines.append(f"  idx {pf.idx:>3}  {pf.type.name.lower():<10} "
                         f"at ({_inches(ph.left):.2f}, {_inches(ph.top):.2f}) "
                         f"{_inches(ph.width):.2f} x {_inches(ph.height):.2f}  {ph.name}")
    return "\n".join(lines)
