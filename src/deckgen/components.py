"""Slide components. Each one draws itself into a Box.

To add a new component type:

    @register("my_type")
    def render_my_type(slide, box: Box, spec: dict, ctx: Context) -> None:
        ...

and use `type: my_type` in the YAML.
"""

from __future__ import annotations

import copy
import math
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

from pptx.chart.data import CategoryChartData, XyChartData
from pptx.chart.plot import LinePlot
from pptx.dml.color import RGBColor
from pptx.enum.chart import (XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION,
                             XL_MARKER_STYLE)
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from . import data as dq
from .layout import Box
from .theme import Theme


@dataclass
class Context:
    theme: Theme
    data: dict[str, list[dict]]
    base_dir: Path


Renderer = Callable[[object, Box, dict, Context], None]
REGISTRY: dict[str, Renderer] = {}


def register(name: str):
    def deco(fn: Renderer) -> Renderer:
        REGISTRY[name] = fn
        return fn
    return deco


def render(slide, box: Box, spec: dict, ctx: Context) -> None:
    """Draws one component. Any component can sit on a `panel` (a tinted backdrop)."""
    kind = spec.get("type")
    if kind not in REGISTRY:
        raise ValueError(f"unknown component type '{kind}'; available: {sorted(REGISTRY)}")
    if panel := spec.get("panel"):
        bg = _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, box,
                    ctx.theme.rgb("light" if panel is True else panel))
        bg.adjustments[0] = min(0.2, 0.12 / min(box.w, box.h))
        box = box.inset(spec.get("padding", 0.2))
    REGISTRY[kind](slide, box, spec, ctx)


# --------------------------------------------------------------------------- helpers

ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
ICON_FONT = "Segoe UI Symbol"


def _emu(box: Box) -> tuple[Emu, Emu, Emu, Emu]:
    return Inches(box.x), Inches(box.y), Inches(box.w), Inches(box.h)


def add_rich_text(paragraph, text: str, theme: Theme, size: float | None,
                  color: str | None = "text", bold: bool = False) -> None:
    """Adds text to a paragraph; **double asterisks** become bold accent-colored runs.

    size/color None leave the run unstyled so it inherits from a template placeholder.
    """
    inherit = size is None
    for i, chunk in enumerate(re.split(r"\*\*(.+?)\*\*", str(text))):
        if not chunk:
            continue
        run = paragraph.add_run()
        run.text = chunk
        emphasized = i % 2 == 1
        if not inherit:
            run.font.name = theme.font
            run.font.size = Pt(size)
        if bold or emphasized:
            run.font.bold = True
        if emphasized:
            run.font.color.rgb = theme.rgb("accent")
        elif color is not None:
            run.font.color.rgb = theme.rgb(color)


def _textbox(slide, box: Box, anchor=MSO_ANCHOR.TOP, margin: float = 0.05):
    tb = slide.shapes.add_textbox(*_emu(box))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    return tf


def _lines(slide, box: Box, lines: list[tuple], theme: Theme, align: str = "left",
           anchor: str = "top") -> None:
    """Textbox with one paragraph per (text, size, color, bold) tuple; empty text is skipped."""
    tf = _textbox(slide, box, anchor=ANCHOR[anchor])
    first = True
    for text, size, color, bold in lines:
        if not text:
            continue
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = ALIGN[align]
        p.space_after = Pt(size * 0.3)
        add_rich_text(p, text, theme, size, color=color, bold=bold)


def _bullet(paragraph, theme: Theme, numbered: bool) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(0.28)))
    pPr.set("indent", str(-Inches(0.28)))
    clr = pPr.makeelement(qn("a:buClr"), {})
    srgb = clr.makeelement(qn("a:srgbClr"), {"val": str(theme.rgb("accent"))})
    clr.append(srgb)
    pPr.append(clr)
    if numbered:
        pPr.append(pPr.makeelement(qn("a:buAutoNum"), {"type": "arabicPeriod"}))
    else:
        pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))


def _fill(shape, color) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _shape(slide, kind, box: Box, color: RGBColor):
    shape = slide.shapes.add_shape(kind, *_emu(box))
    _fill(shape, color)
    shape.shadow.inherit = False
    return shape


def _icon(slide, box: Box, glyph: str, theme: Theme, fill: str = "accent",
          color: str = "FFFFFF") -> None:
    """A glyph (any Unicode symbol) centred in a filled circle."""
    circle = _shape(slide, MSO_SHAPE.OVAL, box, theme.rgb(fill))
    tf = circle.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_rich_text(p, glyph, theme, box.h * 72 * 0.45, color=color, bold=True)
    for run in p.runs:
        if not glyph.isalnum():
            run.font.name = ICON_FONT


def fmt(value, pattern: str | None) -> str:
    if pattern is None:
        if isinstance(value, float):
            return f"{value:,.2f}"
        if isinstance(value, int):
            return f"{value:,}"
        return str(value)
    return pattern.format(value)


def _items(spec: dict, key: str = "title") -> list[dict]:
    """`items` as dicts; a plain string item becomes {key: string}."""
    return [i if isinstance(i, dict) else {key: i} for i in spec.get("items", [])]


def _step_color(i: int, active: int | None) -> str:
    """Colour of step i (1-based) in a sequence where `active` is the current step."""
    if active is None:
        return "primary"
    return "accent" if i == active else ("primary" if i < active else "subtle")


# --------------------------------------------------------------------------- text components

@register("insights")
def render_insights(slide, box: Box, spec: dict, ctx: Context) -> None:
    """Bullet list of takeaways. style: bullets (default) | numbered | callout."""
    theme = ctx.theme
    style = spec.get("style", "bullets")
    size = spec.get("font_size", 16)
    inner = box
    if style == "callout":
        _shape(slide, MSO_SHAPE.RECTANGLE, box, theme.rgb("light"))
        _shape(slide, MSO_SHAPE.RECTANGLE, Box(box.x, box.y, 0.08, box.h), theme.rgb("accent"))
        inner = Box(box.x + 0.3, box.y + 0.2, box.w - 0.5, box.h - 0.4)

    tf = _textbox(slide, inner, anchor=MSO_ANCHOR.MIDDLE if style == "callout" else MSO_ANCHOR.TOP)
    first = True
    if heading := spec.get("heading"):
        p = tf.paragraphs[0]
        add_rich_text(p, heading, theme, size + 2, color="primary", bold=True)
        p.space_after = Pt(8)
        first = False
    for item in spec.get("items", []):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(spec.get("spacing", 10))
        _bullet(p, theme, numbered=style == "numbered")
        add_rich_text(p, item, theme, size)


@register("text")
def render_text(slide, box: Box, spec: dict, ctx: Context) -> None:
    tf = _textbox(slide, box, anchor=ANCHOR[spec.get("valign", "top")])
    for i, para in enumerate(str(spec.get("text", "")).strip().split("\n\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN[spec.get("align", "left")]
        p.space_after = Pt(8)
        add_rich_text(p, para.replace("\n", " "), ctx.theme, spec.get("font_size", 16),
                      color=spec.get("color", "text"), bold=spec.get("bold", False))
        if spec.get("italic"):
            for run in p.runs:
                run.font.italic = True


@register("quote")
def render_quote(slide, box: Box, spec: dict, ctx: Context) -> None:
    """A pull quote: large quotation mark, the quote, and an attribution line."""
    theme = ctx.theme
    mark = _textbox(slide, Box(box.x - 0.1, box.y + box.h / 2 - 2.3, 1.4, 2.6))
    add_rich_text(mark.paragraphs[0], "“", theme, 150, color="accent", bold=True)
    mark.paragraphs[0].runs[0].font.name = "Georgia"
    tf = _textbox(slide, Box(box.x + 1.2, box.y, box.w - 1.2, box.h),
                  anchor=ANCHOR[spec.get("valign", "middle")])
    p = tf.paragraphs[0]
    p.space_after = Pt(18)
    add_rich_text(p, spec.get("text", ""), theme, spec.get("font_size", 30),
                  color=spec.get("color", "primary"))
    for run in p.runs:
        run.font.italic = True
    for line, size, color, bold in ((spec.get("author"), 18, "text", True),
                                    (spec.get("role"), 14, "muted", False)):
        if line:
            p = tf.add_paragraph()
            add_rich_text(p, line, theme, size, color=color, bold=bold)


# --------------------------------------------------------------------------- cards

KPI_STYLES = ("flat", "accent", "outline", "dark")


@register("kpis")
def render_kpis(slide, box: Box, spec: dict, ctx: Context) -> None:
    """Row of metric cards. `value`/`delta` may be literals or computed metric specs.

    style: flat (default) | accent (bar down the left edge) | outline | dark (filled with
    the primary colour). Each item may add an `icon`: any Unicode glyph, shown in a circle.
    """
    theme = ctx.theme
    style = spec.get("style", "flat")
    if style not in KPI_STYLES:
        raise ValueError(f"unknown kpis style '{style}'; valid: {list(KPI_STYLES)}")
    items = spec.get("items", [])
    cards = box.split_cols([1] * len(items), gap=spec.get("gap", 0.25))
    for item, card in zip(items, cards):
        fill = {"dark": "primary", "outline": theme.mix("light", "FFFFFF", 0.7)}.get(style, "light")
        shape = MSO_SHAPE.RECTANGLE if style == "accent" else MSO_SHAPE.ROUNDED_RECTANGLE
        bg = _shape(slide, shape, card, theme.rgb(fill))
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            bg.adjustments[0] = min(0.3, 0.14 / min(card.w, card.h))
        if style == "outline":
            bg.line.color.rgb = theme.rgb("subtle")
            bg.line.width = Pt(1)
        inset = 0.2
        if style == "accent":
            _shape(slide, MSO_SHAPE.RECTANGLE, Box(card.x, card.y, 0.08, card.h), theme.rgb("accent"))
            inset = 0.3
        text_w = card.w - inset - 0.2
        if glyph := item.get("icon"):
            d = min(0.55, card.h * 0.38)
            _icon(slide, Box(card.x + card.w - d - 0.18, card.y + 0.18, d, d), str(glyph), theme)
            text_w -= d + 0.05

        dark = style == "dark"
        label_color = theme.mix("primary", "FFFFFF", 0.7) if dark else "muted"
        value = _resolve_metric(item.get("value"), ctx)
        delta = _resolve_metric(item.get("delta"), ctx)
        tf = _textbox(slide, Box(card.x + inset, card.y + 0.1, text_w, card.h - 0.2),
                      anchor=MSO_ANCHOR.MIDDLE)
        p = tf.paragraphs[0]
        add_rich_text(p, item.get("label", ""), theme, 12, color=label_color)
        p = tf.add_paragraph()
        add_rich_text(p, fmt(value, item.get("format")), theme,
                      item.get("value_size", spec.get("value_size", 30)),
                      color="FFFFFF" if dark else "primary", bold=True)
        if delta is not None:
            text = fmt(delta, item.get("delta_format"))
            is_up = not str(text).lstrip().startswith("-")
            good = is_up if item.get("higher_is_better", True) else not is_up
            p = tf.add_paragraph()
            add_rich_text(p, ("▲ " if is_up else "▼ ") + text.lstrip("+-") +
                          (f"  {item['delta_label']}" if item.get("delta_label") else ""),
                          theme, 12, color="positive" if good else "negative")
        if note := item.get("note"):
            p = tf.add_paragraph()
            add_rich_text(p, note, theme, 12, color=label_color)


def _resolve_metric(value, ctx: Context):
    if isinstance(value, dict):
        rows = _source_rows(value["source"], ctx)
        return dq.metric(rows, value)
    return value


@register("progress")
def render_progress(slide, box: Box, spec: dict, ctx: Context) -> None:
    """Horizontal progress bars: label | bar (value out of `max`) | value, one row per item.

    Each item: label, value (literal or metric), max (default 1), format, target (a marker
    line), color. `highlight: true` draws that bar in the accent colour.
    """
    theme = ctx.theme
    items = spec.get("items", [])
    rows = box.split_rows([1] * len(items), gap=spec.get("gap", 0.1))
    label_w = box.w * spec.get("label_width", 0.3)
    value_w = spec.get("value_width", 1.2)
    size = spec.get("font_size", 15)
    track_color = theme.rgb(theme.mix("light", "subtle", 0.45))
    for item, row in zip(items, rows):
        value = _resolve_metric(item.get("value"), ctx)
        top = _resolve_metric(item.get("max", spec.get("max", 1)), ctx)
        bar_h = min(spec.get("bar_height", 0.3), row.h * 0.45)
        track = Box(row.x + label_w + 0.15, row.y + (row.h - bar_h) / 2,
                    row.w - label_w - value_w - 0.3, bar_h)
        _lines(slide, Box(row.x, row.y, label_w, row.h), [(item.get("label", ""), size, "text", False)],
               theme, anchor="middle")
        bar = _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, track, track_color)
        bar.adjustments[0] = 0.5
        frac = max(0.0, min(1.0, value / top)) if top else 0
        if frac > 0:
            color = item.get("color") or ("accent" if item.get("highlight") else spec.get("color", "primary"))
            fill = _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE,
                          Box(track.x, track.y, max(track.w * frac, bar_h), bar_h), theme.rgb(color))
            fill.adjustments[0] = 0.5
        if (target := item.get("target")) is not None:
            tx = track.x + track.w * min(1.0, _resolve_metric(target, ctx) / top)
            _shape(slide, MSO_SHAPE.RECTANGLE, Box(tx - 0.02, track.y - 0.1, 0.04, bar_h + 0.2),
                   theme.rgb("text"))
        _lines(slide, Box(track.x + track.w + 0.15, row.y, value_w, row.h),
               [(fmt(value, item.get("format", spec.get("format"))), size + 2, "primary", True)],
               theme, align="right", anchor="middle")


# --------------------------------------------------------------------------- sequences

@register("agenda")
def render_agenda(slide, box: Box, spec: dict, ctx: Context) -> None:
    """Numbered chapters. `active: n` highlights chapter n and dims the others, so the same
    agenda can be repeated as a progress tracker. layout: rows (default) | columns."""
    theme = ctx.theme
    items = _items(spec)
    active = spec.get("active")
    columns = spec.get("layout", "rows") == "columns"
    cells = (box.split_cols if columns else box.split_rows)([1] * len(items),
                                                            gap=spec.get("gap", 0.15))
    for n, (item, cell) in enumerate(zip(items, cells), start=1):
        on, dim = active == n, active is not None and active != n
        ink = "subtle" if dim else "primary"
        if on:
            _shape(slide, MSO_SHAPE.RECTANGLE, cell, theme.rgb("light"))
            edge = Box(cell.x, cell.y, cell.w, 0.07) if columns else Box(cell.x, cell.y, 0.08, cell.h)
            _shape(slide, MSO_SHAPE.RECTANGLE, edge, theme.rgb("accent"))
        elif columns:
            _shape(slide, MSO_SHAPE.RECTANGLE, Box(cell.x, cell.y, cell.w, 0.03), theme.rgb("subtle"))
        elif n < len(items):
            _shape(slide, MSO_SHAPE.RECTANGLE, Box(cell.x, cell.y + cell.h + 0.07, cell.w, 0.015),
                   theme.rgb("subtle"))
        number = item.get("number", f"{n:02d}")
        num_color = "accent" if on or active is None else ink
        lines = [(item.get("title", ""), spec.get("font_size", 20), ink, True),
                 (item.get("text", ""), spec.get("font_size", 20) - 6, "subtle" if dim else "muted", False)]
        if columns:
            inner = cell.inset(0.25, 0.3)
            _lines(slide, inner, [(number, spec.get("number_size", 44), num_color, True)] + lines, theme)
        else:
            num_size = spec.get("number_size", min(40, cell.h * 72 * 0.55))
            _lines(slide, Box(cell.x + 0.3, cell.y, 1.2, cell.h),
                   [(number, num_size, num_color, True)], theme, anchor="middle")
            _lines(slide, Box(cell.x + 1.5, cell.y, cell.w - 1.7, cell.h), lines, theme,
                   anchor="middle")


@register("timeline")
def render_timeline(slide, box: Box, spec: dict, ctx: Context) -> None:
    """Steps left to right. style: chevron (default) | dots. Each item: label (inside the
    chevron / above the dot), title and text (below). `active: n` marks the current step:
    earlier steps are drawn as done, later ones greyed."""
    theme = ctx.theme
    items = _items(spec, key="label")
    active = spec.get("active")
    n = len(items)
    if spec.get("style", "chevron") == "chevron":
        band_h = spec.get("band_height", min(0.8, box.h * 0.35))
        depth, gap = band_h * 0.5, 0.06
        step_w = (box.w + (n - 1) * (depth - gap)) / n
        for i, item in enumerate(items, start=1):
            x = box.x + (i - 1) * (step_w - depth + gap)
            color = _step_color(i, active)
            kind = MSO_SHAPE.PENTAGON if i == 1 else MSO_SHAPE.CHEVRON
            shape = _shape(slide, kind, Box(x, box.y, step_w, band_h), theme.rgb(color))
            tf = shape.text_frame
            tf.margin_left, tf.margin_right = Inches(depth + 0.05), Inches(depth * 0.6)
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            add_rich_text(p, item.get("label", ""), theme, spec.get("label_size", 15),
                          color="primary" if color == "subtle" else "FFFFFF", bold=True)
            _lines(slide, Box(x + depth * 0.5, box.y + band_h + 0.2, step_w - depth - 0.05,
                              box.h - band_h - 0.2),
                   [(item.get("title", ""), spec.get("font_size", 16), "primary", True),
                    (item.get("text", ""), spec.get("font_size", 16) - 3, "muted", False)], theme)
        return

    cells = box.split_cols([1] * n, gap=0.2)
    line_y = box.y + 0.75
    first, last = cells[0], cells[-1]
    _shape(slide, MSO_SHAPE.RECTANGLE,
           Box(first.x + first.w / 2, line_y - 0.02, last.x - first.x, 0.04), theme.rgb("subtle"))
    for i, (item, cell) in enumerate(zip(items, cells), start=1):
        color = _step_color(i, active)
        d = 0.42 if i == active else 0.3
        cx = cell.x + cell.w / 2
        if i == active:
            halo = _shape(slide, MSO_SHAPE.OVAL, Box(cx - 0.33, line_y - 0.33, 0.66, 0.66),
                          theme.rgb(theme.mix("accent", "FFFFFF", 0.75)))
            halo.shadow.inherit = False
        _shape(slide, MSO_SHAPE.OVAL, Box(cx - d / 2, line_y - d / 2, d, d), theme.rgb(color))
        _lines(slide, Box(cell.x, box.y, cell.w, 0.45),
               [(item.get("label", ""), 13, "accent" if i == active else "muted", True)],
               theme, align="center", anchor="bottom")
        _lines(slide, Box(cell.x, line_y + 0.4, cell.w, box.h - 1.15),
               [(item.get("title", ""), spec.get("font_size", 16), "primary", True),
                (item.get("text", ""), spec.get("font_size", 16) - 3, "muted", False)],
               theme, align="center")


@register("cycle")
def render_cycle(slide, box: Box, spec: dict, ctx: Context) -> None:
    """A lifecycle: steps around a ring, clockwise from the top, with arrows between them.

    Each item: title, text, and optional `icon` (a glyph shown in the node instead of its
    number). `center` puts text in the middle; `active: n` highlights step n.
    """
    theme = ctx.theme
    items = _items(spec)
    active = spec.get("active")
    n = len(items)
    cx, cy = box.x + box.w / 2, box.y + box.h / 2
    node = spec.get("node_size", min(0.95, box.h * 0.19))
    label_h = spec.get("label_height", 0.8)
    r = min(box.h / 2 - node / 2 - 0.15 - label_h, box.w * spec.get("radius", 0.2))
    ring = _shape(slide, MSO_SHAPE.DONUT, Box(cx - r, cy - r, 2 * r, 2 * r),
                  theme.rgb(theme.mix("light", "subtle", 0.6)))
    ring.adjustments[0] = 0.07 / (2 * r)

    def at(deg: float, dist: float) -> tuple[float, float]:
        rad = math.radians(deg)
        return cx + dist * math.cos(rad), cy + dist * math.sin(rad)

    step = 360 / n
    for i in range(n):  # arrowheads halfway between nodes, pointing clockwise
        deg = -90 + step * (i + 0.5)
        ax, ay = at(deg, r)
        arrow = _shape(slide, MSO_SHAPE.ISOSCELES_TRIANGLE, Box(ax - 0.12, ay - 0.1, 0.24, 0.2),
                       theme.rgb(theme.mix("light", "subtle", 0.9)))
        arrow.rotation = deg + 180

    label_w = spec.get("label_width", min(2.8, box.w / 2 - r - node / 2 - 0.25))
    for i, item in enumerate(items, start=1):
        deg = -90 + step * (i - 1)
        px, py = at(deg, r)
        color = "accent" if i == active else "primary"
        _icon(slide, Box(px - node / 2, py - node / 2, node, node), str(item.get("icon", i)),
              theme, fill=color)
        lx, ly = at(deg, r + node / 2 + 0.15)
        cos, sin = math.cos(math.radians(deg)), math.sin(math.radians(deg))
        h = label_h if abs(cos) <= 0.35 else 1.1
        if cos > 0.35:
            lbox, align = Box(lx, ly - h / 2, label_w, h), "left"
        elif cos < -0.35:
            lbox, align = Box(lx - label_w, ly - h / 2, label_w, h), "right"
        elif sin < 0:
            lbox, align = Box(lx - label_w / 2, ly - h, label_w, h), "center"
        else:
            lbox, align = Box(lx - label_w / 2, ly, label_w, h), "center"
        anchor = "middle" if abs(cos) > 0.35 else ("bottom" if sin < 0 else "top")
        _lines(slide, lbox, [(item.get("title", ""), spec.get("font_size", 16),
                              "accent" if i == active else "primary", True),
                             (item.get("text", ""), spec.get("font_size", 16) - 3, "muted", False)],
               theme, align=align, anchor=anchor)
    if center := spec.get("center"):
        inner = r - node / 2 - 0.1
        _lines(slide, Box(cx - inner, cy - inner, 2 * inner, 2 * inner),
               [(line, spec.get("center_size", 20), "primary", True)
                for line in str(center).split("\n")], theme, align="center", anchor="middle")


# --------------------------------------------------------------------------- data components

def _source_rows(name: str, ctx: Context) -> list[dict]:
    if name not in ctx.data:
        raise ValueError(f"unknown data source '{name}'; defined: {sorted(ctx.data)}")
    return ctx.data[name]


def _series_label(spec: dict, column: str) -> str:
    return spec.get("series_names", {}).get(column, column.replace("_", " ").title())


def _chart_series(spec: dict, ctx: Context) -> tuple[list, dict[str, list]]:
    """Returns (categories, {series_name: values}) from inline data or a named source."""
    src = spec.get("data")
    if isinstance(src, dict):  # inline: {categories: [...], series: {name: [...]}}
        return src["categories"], dict(src["series"])

    rows = dq.query(_source_rows(src, ctx), spec)
    x = spec["x"]
    ys = spec["y"] if isinstance(spec["y"], list) else [spec["y"]]
    if series_by := spec.get("series_by"):
        dq.check_columns(rows, [x, ys[0], series_by], src)
        return dq.pivot(rows, x, ys[0], series_by)
    dq.check_columns(rows, [x, *ys], src)
    return [r[x] for r in rows], {_series_label(spec, y): [r[y] for r in rows] for y in ys}


def _waterfall(categories: list, values: list, totals: list) -> dict[str, list]:
    """Splits a bridge into stacked series; `totals` are categories drawn as full bars."""
    series = {"Base": [], "Total": [], "Increase": [], "Decrease": []}
    running = 0
    for cat, v in zip(categories, values):
        if cat in totals:
            running = v
            row = (0, v, 0, 0)
        else:
            start, running = running, running + v
            row = (min(start, running), 0, max(v, 0), max(-v, 0))
        for name, part in zip(series, row):
            series[name].append(part)
    return series


def _move_to_line_plot(chart, names: set[str]) -> None:
    """Moves the named series of a bar chart into a line plot on the same axes (a combo)."""
    bar = chart._chartSpace.chart.plotArea.find(qn("c:barChart"))
    line = bar.makeelement(qn("c:lineChart"), {})
    for tag, val in (("c:grouping", "standard"), ("c:varyColors", "0")):
        line.append(line.makeelement(qn(tag), {"val": val}))
    for s in list(chart.plots[0].series):
        if s.name not in names:
            continue
        ser = s._element
        bar.remove(ser)
        if (inv := ser.find(qn("c:invertIfNegative"))) is not None:
            ser.remove(inv)
        ser.find(qn("c:cat")).addprevious(ser.makeelement(qn("c:marker"), {}))
        ser.append(ser.makeelement(qn("c:smooth"), {"val": "0"}))
        line.append(ser)
    line.append(line.makeelement(qn("c:marker"), {"val": "1"}))
    for ax in bar.findall(qn("c:axId")):
        line.append(copy.deepcopy(ax))
    bar.addnext(line)


CHART_TYPES = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "stacked_column": XL_CHART_TYPE.COLUMN_STACKED,
    "stacked_column_100": XL_CHART_TYPE.COLUMN_STACKED_100,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "stacked_bar": XL_CHART_TYPE.BAR_STACKED,
    "stacked_bar_100": XL_CHART_TYPE.BAR_STACKED_100,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "area": XL_CHART_TYPE.AREA,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
    "scatter": XL_CHART_TYPE.XY_SCATTER,
    "combo": XL_CHART_TYPE.COLUMN_CLUSTERED,  # columns plus the `line:` series as lines
    "waterfall": XL_CHART_TYPE.COLUMN_STACKED,
}
BAR_KINDS = {"column", "stacked_column", "stacked_column_100", "bar", "stacked_bar",
             "stacked_bar_100", "combo", "waterfall"}
LEGEND = {"bottom": XL_LEGEND_POSITION.BOTTOM, "right": XL_LEGEND_POSITION.RIGHT,
          "top": XL_LEGEND_POSITION.TOP, "left": XL_LEGEND_POSITION.LEFT}


def _label_last_point(series, n_points: int, text: str, color: RGBColor, size: float) -> None:
    label = series.points[n_points - 1].data_label
    label.text_frame.text = text
    run = label.text_frame.paragraphs[0].runs[0]
    run.font.size, run.font.bold = Pt(size), True
    run.font.color.rgb = color
    label.position = XL_LABEL_POSITION.RIGHT


def _text_on(fill: RGBColor, theme: Theme) -> RGBColor:
    """White text on a dark fill, the theme's text colour on a light one."""
    return theme.rgb("FFFFFF" if theme.is_dark(str(fill)) else "text")


def _label_point_on_fill(point, fill: RGBColor, theme: Theme, number_format: str,
                         percentage: bool) -> None:
    """Styles one pie/doughnut slice's label so it stays readable on the slice colour."""
    label = point.data_label
    label.font.size = Pt(10)
    label.font.color.rgb = _text_on(fill, theme)
    dlbl = label._dLbl
    dlbl.find(qn("c:showVal")).set("val", "0" if percentage else "1")
    dlbl.find(qn("c:showPercent")).set("val", "1" if percentage else "0")
    num_fmt = dlbl.makeelement(qn("c:numFmt"), {"formatCode": number_format, "sourceLinked": "0"})
    dlbl.find(qn("c:txPr")).addprevious(num_fmt)


def _reserve_right_margin(chart, inches: float, chart_width: float) -> None:
    """Shrinks the plot area so labels placed right of the last point sit outside the plot."""
    plot_area = chart._chartSpace.chart.plotArea
    for old in plot_area.findall(qn("c:layout")):
        plot_area.remove(old)
    left, top, bottom = 0.09, 0.04, 0.10
    width = 1 - left - min(0.4, inches / chart_width)
    layout = plot_area.makeelement(qn("c:layout"), {})
    manual = layout.makeelement(qn("c:manualLayout"), {})
    for tag, val in (("layoutTarget", "inner"), ("xMode", "edge"), ("yMode", "edge"),
                     ("x", left), ("y", top), ("w", width), ("h", 1 - top - bottom)):
        manual.append(manual.makeelement(qn(f"c:{tag}"), {"val": str(val)}))
    layout.append(manual)
    plot_area.insert(0, layout)


@register("chart")
def render_chart(slide, box: Box, spec: dict, ctx: Context) -> None:
    """Native (editable) PowerPoint chart. kind: see CHART_TYPES.

    Emphasis: `highlight` (category or series names) draws those in the accent colour and
    the rest in grey; `direct_labels` names each line at its end instead of a legend;
    `reference: {value, label}` adds a dashed reference line.
    """
    theme = ctx.theme
    kind = spec.get("kind", "column")
    if kind not in CHART_TYPES:
        raise ValueError(f"unknown chart kind '{kind}'; available: {sorted(CHART_TYPES)}")
    number_format = spec.get("number_format", "General")
    highlight = spec.get("highlight") or []
    highlight = set(highlight if isinstance(highlight, list) else [highlight])
    reference = spec.get("reference")
    ref_name = reference.get("label", "Target") if reference else None
    if reference and kind not in {"column", "stacked_column", "line", "area", "combo"}:
        raise ValueError("`reference` works with column, stacked_column, line, area and combo")

    categories: list = []
    if kind == "scatter":
        rows = dq.query(_source_rows(spec["data"], ctx), spec)
        dq.check_columns(rows, [spec["x"], spec["y"]], spec["data"])
        cd = XyChartData()
        groups = {}
        for r in rows:
            groups.setdefault(r[spec["series_by"]] if spec.get("series_by") else spec["y"], []).append(r)
        for name, members in groups.items():
            s = cd.add_series(str(name), number_format=number_format)
            for r in members:
                s.add_data_point(r[spec["x"]], r[spec["y"]])
    else:
        categories, series = _chart_series(spec, ctx)
        if kind == "waterfall":
            series = _waterfall(categories, next(iter(series.values())), spec.get("totals", []))
        if reference:
            if kind in BAR_KINDS:  # an empty last slot gives the line's end label room beside the bars
                categories = [*categories, ""]
                series = {name: [*values, None] for name, values in series.items()}
            series[ref_name] = [reference["value"]] * len(categories)
        cd = CategoryChartData(number_format=number_format)
        cd.categories = categories
        for name, values in series.items():
            cd.add_series(name, values)

    chart = slide.shapes.add_chart(CHART_TYPES[kind], *_emu(box), cd).chart
    chart.font.name = theme.font
    chart.font.size = Pt(spec.get("font_size", 11))
    chart.font.color.rgb = theme.rgb("text")

    line_series = set()
    if kind == "combo":
        line_series = {_series_label(spec, c) for c in spec.get("line", [])} | set(spec.get("line", []))
    if ref_name and kind in BAR_KINDS:
        line_series.add(ref_name)
    if line_series:
        _move_to_line_plot(chart, line_series)

    if title := spec.get("title"):
        chart.has_title = True
        chart.chart_title.text_frame.text = title
        run = chart.chart_title.text_frame.paragraphs[0].runs[0]
        run.font.size, run.font.bold = Pt(13), True
        run.font.color.rgb = theme.rgb("primary")
    else:
        chart.has_title = False

    plots = list(chart.plots)
    all_series = [(s, isinstance(p, LinePlot)) for p in plots for s in p.series]
    shown = [s for s, _ in all_series if s.name not in (ref_name, "Base")]
    direct = spec.get("direct_labels", False)
    default_legend = ("bottom" if (len(shown) > 1 or kind in ("pie", "doughnut"))
                      and not direct and kind != "waterfall" else "none")
    legend = spec.get("legend", default_legend)
    chart.has_legend = legend != "none"
    if chart.has_legend:
        chart.legend.position = LEGEND[legend]
        chart.legend.include_in_layout = False

    end_labels = [s.name for s in shown] if direct and kind in ("line", "combo") else []
    if ref_name:
        end_labels.append(ref_name)
    if end_labels and kind not in ("pie", "doughnut", "scatter"):
        longest = max(len(name) for name in end_labels)
        _reserve_right_margin(chart, longest * 0.09 + 0.2, box.w)

    plot = plots[0]
    fills: list[RGBColor] = []  # fill of each pie slice, or of each series, for label contrast
    if kind in ("pie", "doughnut"):
        for i, point in enumerate(plot.series[0].points):
            point.format.fill.solid()
            fills.append(
                (theme.rgb("accent") if categories[i] in highlight else theme.rgb("subtle"))
                if highlight else theme.series_color(i))
            point.format.fill.fore_color.rgb = fills[-1]
    else:
        waterfall_colors = {"Total": "primary", "Increase": "positive", "Decrease": "negative"}
        for i, (s, as_line) in enumerate(all_series):
            color = theme.series_color(i)
            if kind == "waterfall":
                if s.name == "Base":
                    s.format.fill.background()
                    s.format.line.fill.background()
                    continue
                color = theme.rgb(waterfall_colors[s.name])
            elif highlight and len(shown) > 1:
                color = theme.rgb("accent" if s.name in highlight else "subtle")
            if s.name == ref_name:
                s.format.line.color.rgb = theme.rgb("muted")
                s.format.line.width = Pt(1.5)
                s.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
                s.marker.style = XL_MARKER_STYLE.NONE
                s.marker.format.fill.background()
                s.marker.format.line.fill.background()
                s.smooth = False
                _label_last_point(s, len(categories), ref_name, theme.rgb("muted"), 11)
            elif kind in ("line", "scatter") or as_line:
                s.format.line.color.rgb = color
                s.format.line.width = Pt(3 if s.name in highlight else 2.25)
                s.marker.format.fill.solid()
                s.marker.format.fill.fore_color.rgb = color
                s.marker.format.line.color.rgb = color
                s.smooth = False
                if kind == "scatter":
                    s.format.line.fill.background()
                    s.marker.size = 10
                elif direct:
                    _label_last_point(s, len(categories), s.name, color, 12)
            else:
                s.format.fill.solid()
                s.format.fill.fore_color.rgb = color
                fills.append(color)
                if highlight and len(shown) == 1:
                    for j, cat in enumerate(categories):
                        point = s.points[j]
                        point.format.fill.solid()
                        point.format.fill.fore_color.rgb = theme.rgb(
                            "accent" if cat in highlight else "subtle")
        if kind in BAR_KINDS:
            plot.gap_width = spec.get("gap_width", 60)
            if kind.startswith("stacked") or kind == "waterfall":
                plot.overlap = 100
        va = chart.value_axis
        va.has_major_gridlines = spec.get("gridlines", True)
        if va.has_major_gridlines:
            va.major_gridlines.format.line.color.rgb = theme.rgb(theme.mix("light", "subtle", 0.5))
        va.format.line.fill.background()
        va.tick_labels.number_format = spec.get(
            "axis_format", "0%" if kind.endswith("_100") else number_format)
        va.tick_labels.number_format_is_linked = False
        if "axis_min" in spec:
            va.minimum_scale = spec["axis_min"]
        if "axis_max" in spec:
            va.maximum_scale = spec["axis_max"]
        ca = chart.category_axis
        ca.format.line.color.rgb = theme.rgb("muted")
        ca.has_major_gridlines = False
        if kind in ("bar", "stacked_bar", "stacked_bar_100"):
            ca.reverse_order = True  # keep first row at the top
            va._element.find(qn("c:crosses")).set("val", "max")  # ...and the value axis at the bottom

    label_format = spec.get("label_format", number_format)
    if kind == "waterfall" and spec.get("data_labels", True):
        for s in plot.series:
            if s.name == "Base":
                continue
            sign = {"Increase": "+", "Decrease": "-"}.get(s.name, "")
            dl = s.data_labels
            dl.show_value = True
            dl.number_format = f"{sign}{label_format};{sign}{label_format};;"
            dl.number_format_is_linked = False
            dl.font.size = Pt(10)
            dl.font.bold = True
            dl.font.color.rgb = _text_on(theme.rgb(waterfall_colors[s.name]), theme)
    elif spec.get("data_labels"):
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format = label_format
        dl.number_format_is_linked = False
        dl.font.size = Pt(10)
        if kind in ("pie", "doughnut"):
            dl.show_percentage = spec.get("show_percentage", False)
            dl.show_value = not dl.show_percentage
            for point, fill in zip(plot.series[0].points, fills):
                _label_point_on_fill(point, fill, theme, label_format, dl.show_percentage)
        elif kind.startswith("stacked"):
            for series, fill in zip(plot.series, fills):
                series_dl = series.data_labels
                series_dl.show_value = True
                series_dl.number_format = label_format
                series_dl.number_format_is_linked = False
                series_dl.font.size = Pt(10)
                series_dl.font.color.rgb = _text_on(fill, theme)


@register("table")
def render_table(slide, box: Box, spec: dict, ctx: Context) -> None:
    theme = ctx.theme
    rows = dq.query(_source_rows(spec["data"], ctx), spec)
    columns = spec.get("columns") or list(rows[0])
    if isinstance(columns, list):
        columns = {c: c.replace("_", " ").title() for c in columns}
    dq.check_columns(rows, list(columns), spec["data"])
    formats = spec.get("formats", {})
    size = spec.get("font_size", 12)

    row_h = min(0.45, box.h / (len(rows) + 1))
    shape = slide.shapes.add_table(len(rows) + 1, len(columns), Inches(box.x), Inches(box.y),
                                   Inches(box.w), Inches(row_h * (len(rows) + 1)))
    table = shape.table
    highlight = spec.get("highlight_top", 0)

    def write(cell, text, *, header=False, numeric=False, row_i=0):
        cell.text = ""
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT if numeric else PP_ALIGN.LEFT
        add_rich_text(p, text, theme, size, color="FFFFFF" if header else "text",
                      bold=header or row_i < highlight)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = cell.margin_right = Inches(0.1)
        cell.fill.solid()
        if header:
            cell.fill.fore_color.rgb = theme.rgb("primary")
        else:
            cell.fill.fore_color.rgb = theme.rgb("light") if row_i % 2 else theme.rgb("FFFFFF")

    numeric_cols = {c for c in columns if isinstance(rows[0][c], (int, float))}
    for j, (col, header) in enumerate(columns.items()):
        write(table.cell(0, j), header, header=True, numeric=col in numeric_cols)
    for i, row in enumerate(rows, start=1):
        for j, col in enumerate(columns):
            write(table.cell(i, j), fmt(row[col], formats.get(col)),
                  numeric=col in numeric_cols, row_i=i - 1)


@register("image")
def render_image(slide, box: Box, spec: dict, ctx: Context) -> None:
    path = (ctx.base_dir / spec["path"]).resolve()
    if not path.exists():
        raise FileNotFoundError(f"image not found: {path}")
    slide.shapes.add_picture(str(path), Inches(box.x), Inches(box.y), height=Inches(box.h))
